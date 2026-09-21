"""Read-only dataset audit; no patient demographics or inferred image labels."""
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from .dicom import read_dicom

LABEL_COLUMNS = {
    'spine_position': 2, 'spine_axis': 3, 'spine_artifact': 4,
    'right_hip_position': 5, 'right_hip_roi': 6,
    'left_hip_position': 7, 'left_hip_roi': 8,
    'spine_quality': 9, 'right_hip_quality': 10, 'left_hip_quality': 11,
}
GROUPS = {
    'spine': ('spine_position', 'spine_axis', 'spine_artifact'),
    'right_hip': ('right_hip_position', 'right_hip_roi'),
    'left_hip': ('left_hip_position', 'left_hip_roi'),
}


def audit_labels(path: str | Path) -> dict[str, Any]:
    """Audit the supplied two-header Excel layout. 1 means violation.

    Empty cells remain None with an explicit false mask. Free-text comments are
    intentionally not exported. Unknown layouts fail instead of guessing.
    """
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = book.active
        rows = sheet.iter_rows(values_only=True)
        header = next(rows, ())
        subheader = next(rows, ())
        if len(header) < 12 or str(header[1]).strip().lower() != 'study':
            raise ValueError('Expected the supplied two-header Excel schema (study in B1)')
        expected = {2: 'укладка', 3: 'ось', 4: 'предмет', 5: 'позиционирование',
                    6: 'области', 7: 'позиционирование', 8: 'области',
                    9: 'Позвоночник', 10: 'правого', 11: 'левого'}
        if any(len(subheader) <= i or token.lower() not in str(subheader[i]).lower() for i, token in expected.items()):
            raise ValueError('Unrecognized label column layout')
        records, invalid, contradictions = [], [], []
        for excel_row, row in enumerate(rows, 3):
            if len(row) < 2 or row[1] is None:
                continue
            labels = {}
            for name, index in LABEL_COLUMNS.items():
                value = row[index] if index < len(row) else None
                if value is None or (isinstance(value, str) and not value.strip()):
                    labels[name] = None
                elif value in (0, 1, '0', '1'):
                    labels[name] = int(value)
                else:
                    labels[name] = None
                    invalid.append({'excel_row': excel_row, 'label': name, 'reason': 'Expected binary 0/1'})
            record = {'excel_row': excel_row, 'study_folder': str(row[1]).strip(),
                      'labels': labels, 'label_mask': {k: v is not None for k, v in labels.items()}}
            records.append(record)
            for region, names in GROUPS.items():
                quality = labels[f'{region}_quality']
                parts = [labels[k] for k in names]
                if quality is not None and all(v is not None for v in parts) and quality != int(any(parts)):
                    contradictions.append({'excel_row': excel_row, 'region': region,
                                           'quality': quality, 'violations': parts})
        counts = {k: {'observed': sum(r['labels'][k] is not None for r in records),
                      'positive': sum(r['labels'][k] == 1 for r in records),
                      'missing': sum(r['labels'][k] is None for r in records)} for k in LABEL_COLUMNS}
        folder_counts = Counter(r['study_folder'] for r in records)
        return {'row_count': len(records), 'records': records, 'counts': counts,
                'contradictions': contradictions, 'invalid_labels': invalid,
                'duplicate_study_rows': sorted(k for k, count in folder_counts.items() if count > 1)}
    finally:
        book.close()


def audit_dataset(root: str | Path, labels_path: str | Path | None = None) -> dict[str, Any]:
    root = Path(root)
    if not root.is_dir():
        raise ValueError('Dataset root must be an existing directory')
    labels = audit_labels(labels_path) if labels_path is not None else None
    known_folders = {r['study_folder'] for r in labels['records']} if labels else set()
    manifest, failures = [], []
    hashes, folder_uids, uid_folders = defaultdict(list), defaultdict(set), defaultdict(set)
    files = sorted(p for p in root.rglob('*') if p.is_file() and p.suffix.lower() in {'.dcm', '.dicom'})
    for path in files:
        relative = path.relative_to(root)
        try:
            image = read_dicom(path)
        except ValueError:
            # Detailed decoder errors may contain metadata; keep audit export minimal.
            failures.append({'path': str(relative), 'reason': 'DICOM decoding failed'})
            continue
        matches = [part for part in relative.parts[:-1] if part in known_folders]
        folder = matches[0] if len(matches) == 1 else None
        record = {'path': str(relative), 'study_uid': image.study_uid, 'image_uid': image.image_uid,
                  'pixel_hash': image.pixel_hash, 'shape': list(image.pixels.shape),
                  'spacing_mm': image.spacing_mm, 'warnings': image.warnings, 'label_study_folder': folder}
        manifest.append(record)
        hashes[image.pixel_hash].append(record)
        if folder:
            folder_uids[folder].add(image.study_uid)
            uid_folders[image.study_uid].add(folder)
    duplicates = [{'pixel_hash': h, 'paths': [r['path'] for r in records],
                   'study_uids': sorted({r['study_uid'] for r in records})}
                  for h, records in hashes.items() if len(records) > 1]
    image_counts = Counter(r['image_uid'] for r in manifest)
    return {'file_count': len(files), 'decoded_count': len(manifest), 'failed_count': len(failures),
            'study_count': len({r['study_uid'] for r in manifest}), 'unique_pixel_count': len(hashes),
            'duplicate_file_count': len(manifest) - len(hashes), 'duplicate_groups': duplicates,
            'cross_study_duplicate_groups': [g for g in duplicates if len(g['study_uids']) > 1],
            'duplicate_image_uids': sorted(k for k, count in image_counts.items() if count > 1),
            'manifest': manifest, 'failures': failures, 'labels': labels,
            'folder_to_study_uid': {k: sorted(v) for k, v in sorted(folder_uids.items())},
            'ambiguous_folder_mappings': sorted(k for k, v in folder_uids.items() if len(v) > 1),
            'ambiguous_uid_mappings': sorted(k for k, v in uid_folders.items() if len(v) > 1),
            'unmatched_label_folders': sorted(known_folders - folder_uids.keys()),
            'image_labels_assigned': False}
